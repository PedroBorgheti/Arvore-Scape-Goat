from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "relatorio_projeto_scapegoat.docx"

NAVY = "17324D"
BLUE = "2B6C8A"
GOLD = "D9A441"
LIGHT_BLUE = "EAF2F6"
LIGHT_GOLD = "FBF4E3"
LIGHT_GRAY = "F4F6F8"
DARK = RGBColor(35, 43, 51)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color="B8C4CC", size="8"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color=DARK, size=9.5, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, header=True, first_column=False):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            set_cell_border(cell)
            if header and row_index == 0:
                set_cell_shading(cell, NAVY)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            if first_column and column_index == 0 and row_index > 0:
                set_cell_shading(cell, LIGHT_BLUE)
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    for run in paragraph.runs:
        run.font.name = "Aptos Display"
        run.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)
    return paragraph


def add_body(document, text, bold_prefix=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)
        run.font.color.rgb = DARK
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(10.5)
    run.font.color.rgb = DARK


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK

    for section in document.sections:
        header = section.header.paragraphs[0]
        header.text = "PROJETO PRÁTICO 1  |  ÁRVORES E GRAFOS"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(BLUE)
        add_page_number(section.footer.paragraphs[0])


def add_cover(document):
    for _ in range(4):
        document.add_paragraph()
    eyebrow = document.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eyebrow.add_run("UNIVERSIDADE FEDERAL DO PARANÁ")
    run.font.name = "Aptos"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    run = title.add_run("SCAPEGOAT TREE")
    run.font.name = "Aptos Display"
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Implementação e avaliação experimental em C")
    run.font.name = "Aptos"
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    line = document.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    run.font.color.rgb = RGBColor.from_string(GOLD)

    document.add_paragraph()
    identification = document.add_table(rows=4, cols=2)
    identification.autofit = False
    labels = ["Disciplina", "Professor", "Integrantes", "Data"]
    values = [
        "Árvores e Grafos",
        "Evandro Pizzini",
        "Pedro Borgheti\nMaísa V. Kuhne\nVinicius Nunes",
        "22 de setembro de 2026",
    ]
    for row, label, value in zip(identification.rows, labels, values):
        row.cells[0].width = Cm(4.0)
        row.cells[1].width = Cm(10.0)
        set_cell_text(row.cells[0], label.upper(), bold=True, color=RGBColor.from_string(NAVY))
        set_cell_text(row.cells[1], value, color=DARK)
        set_cell_shading(row.cells[0], LIGHT_GOLD)
        set_cell_border(row.cells[0], GOLD, "10")
        set_cell_border(row.cells[1], "D5DDE2", "8")
    document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Relatório técnico e resultados do benchmark")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    document.add_page_break()


def add_workload_table(document):
    rows = [
        ("Inserção sequencial", "10.000", "100% put", "Chaves de 0 a 9.999"),
        ("Inserção aleatória", "10.000", "100% put", "Chaves pseudoaleatórias"),
        ("Carga mista", "30.000", "50% put / 40% get / 10% delete", "Preload de 10%"),
        ("Leitura intensiva", "30.000", "90% get / 5% put / 5% delete", "Preload de 10%"),
        ("Escrita intensiva", "30.000", "70% put / 10% get / 20% delete", "Preload de 10%"),
        ("Zipfian hotset", "30.000", "100% get", "Concentração nas chaves iniciais"),
    ]
    table = document.add_table(rows=1, cols=4)
    headers = ["Carga", "Operações", "Distribuição", "Descrição"]
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
    for data in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, data):
            set_cell_text(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if value.isnumeric() else None)
    style_table(table, first_column=True)


def add_results_table(document):
    rows = [
        ("Inserção sequencial", "3,702", "370,2", "10.000", "0"),
        ("Inserção aleatória", "0,771", "77,1", "6.340", "0"),
        ("Carga mista", "3,330", "111,0", "7.137", "6.088"),
        ("Leitura intensiva", "2,677", "89,2", "1.992", "10.286"),
        ("Escrita intensiva", "4,105", "136,8", "7.307", "11.842"),
        ("Zipfian hotset", "5,870", "195,7", "10.000", "41.842"),
    ]
    table = document.add_table(rows=1, cols=5)
    headers = ["Carga", "Tempo (ms)", "ns/op", "n final", "Checksum"]
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
    for data in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, data)):
            set_cell_text(cell, value, bold=index == 0, align=WD_ALIGN_PARAGRAPH.CENTER if index else None)
    style_table(table, first_column=True)


def build_report():
    document = Document()
    configure_document(document)
    add_cover(document)

    add_heading(document, "Sumário", 1)
    for item in [
        "1. Apresentação",
        "2. Fundamentação e objetivo",
        "3. Implementação",
        "4. Metodologia experimental",
        "5. Resultados",
        "6. Discussão e conclusão",
    ]:
        add_bullet(document, item)

    add_heading(document, "1. Apresentação", 1)
    add_body(document, "Este trabalho apresenta uma implementação da Scapegoat Tree, uma estrutura de dados baseada em árvore binária de busca que mantém seu desempenho por meio da reconstrução de subárvores desbalanceadas. O projeto foi desenvolvido em C e acompanhado por um benchmark com cargas de leitura, escrita e operações mistas.")
    add_body(document, "A proposta combina a implementação da estrutura, a explicação das decisões algorítmicas e uma avaliação experimental reproduzível. Os resultados apresentados são uma amostra executada em macOS e devem ser interpretados como uma comparação interna entre cargas, não como uma medida universal de desempenho.")

    add_heading(document, "2. Fundamentação e objetivo", 1)
    add_body(document, "Em uma árvore binária de busca comum, a inserção de chaves em ordem pode produzir uma estrutura semelhante a uma lista encadeada. Nesse caso, operações que normalmente seriam logarítmicas podem se aproximar de O(n). A Scapegoat Tree combate esse problema reconstruindo uma subárvore quando o caminho de inserção fica profundo demais.")
    add_body(document, "O parâmetro alpha controla a tolerância ao desequilíbrio. Neste projeto foi usado alpha = 0,70: valores menores tornam a estrutura mais exigente e podem causar mais reconstruções; valores maiores permitem árvores mais altas e reduzem a frequência dessas operações.")
    add_body(document, "O objetivo experimental é observar como a estrutura se comporta em diferentes perfis de uso, medindo o tempo total, o tempo médio por operação, o tamanho final da árvore e um checksum das buscas realizadas.")

    add_heading(document, "3. Implementação", 1)
    add_body(document, "Cada nó armazena uma chave inteira e ponteiros para as subárvores esquerda e direita. O cabeçalho ScapegoatTree mantém a raiz, a quantidade atual de nós, o maior tamanho alcançado desde a última reconstrução total e o fator alpha.")
    add_body(document, "A inserção começa como uma operação tradicional de árvore binária de busca. Depois que a nova chave é inserida, sua profundidade é comparada com o limite logarítmico calculado a partir de alpha. Quando esse limite é ultrapassado, o algoritmo percorre o caminho de volta e reconstrói o primeiro ancestral que viola a regra de peso.")
    add_body(document, "A reconstrução percorre a subárvore em ordem, guarda os ponteiros dos nós em um vetor ordenado e escolhe recursivamente o elemento central como raiz. Assim, os nós existentes são reaproveitados e a subárvore retorna a uma forma equilibrada. A remoção segue as regras de uma árvore binária de busca e pode provocar uma reconstrução completa quando a quantidade de nós cai em relação ao pico registrado.")

    add_heading(document, "4. Metodologia experimental", 1)
    add_body(document, "O benchmark utiliza clock_gettime(CLOCK_MONOTONIC), medindo somente o período das operações de cada carga. Foram usadas 10.000 chaves nas cargas de inserção e 30.000 operações nas cargas mistas. O gerador pseudoaleatório xorshift usa a semente fixa 0x12345678, o que torna as sequências reproduzíveis.")
    add_body(document, "As cargas foram escolhidas para separar situações de estresse estrutural e perfis de acesso mais próximos de um banco chave-valor. Inserções repetidas de chaves já existentes não aumentam o tamanho da árvore, enquanto as remoções podem reduzir o tamanho e acionar a política de reconstrução.")
    add_workload_table(document)
    add_body(document, "Tabela 1 - Cargas utilizadas no benchmark.")

    add_heading(document, "5. Resultados", 1)
    add_body(document, "A tabela abaixo registra uma execução local. Os valores são apresentados com vírgula decimal para facilitar a leitura no relatório; o arquivo CSV mantém a saída bruta produzida pelo programa.")
    add_results_table(document)
    add_body(document, "Tabela 2 - Resultados observados em macOS. Tempo e ns/op variam conforme processador, compilador e carga do sistema.")

    add_heading(document, "6. Discussão e conclusão", 1)
    add_body(document, "A inserção sequencial é o caso mais agressivo para uma árvore binária de busca sem balanceamento. Mesmo nesse cenário, as reconstruções impedem que a estrutura permaneça degradada. A inserção aleatória apresentou menor tempo total nesta amostra, enquanto as cargas de leitura e escrita revelam o efeito combinado de buscas, inserções, remoções e reconstruções.")
    add_body(document, "A carga mista representa um uso equilibrado, e a leitura intensiva concentra a maior parte do trabalho em buscas. Já a escrita intensiva realiza mais inserções, o que aumenta a quantidade de mudanças estruturais. O hotset concentra as buscas em uma região pequena do espaço de chaves, permitindo observar um padrão de acesso não uniforme.")
    add_body(document, "Conclui-se que a Scapegoat Tree oferece uma alternativa simples para manter uma árvore de busca com bom comportamento assintótico sem armazenar informações de balanceamento em cada nó. Como continuidade, seria possível repetir os testes com diferentes valores de alpha, tamanhos de entrada e compiladores, além de comparar diretamente com uma árvore AVL ou Red-Black Tree.")

    add_heading(document, "Referências do projeto", 1)
    add_bullet(document, "Código da estrutura: Projeto/src/ScapegoatTree.c e Projeto/src/ScapegoatTree.h")
    add_bullet(document, "Benchmark: Projeto/src/main.c")
    add_bullet(document, "Dados brutos: Projeto/results/resultados.csv")
    add_bullet(document, "Execução reproduzível: Projeto/scripts/run.sh ou Dockerfile")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()